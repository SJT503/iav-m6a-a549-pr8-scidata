# ---------------------------------------------------------------------------
# Paths and settings come from config/config.ini; nothing is machine-specific.
# ---------------------------------------------------------------------------
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _config import cfg

# -*- coding: utf-8 -*-
"""Build Tables 1-4 for the revised manuscript in the format of the original Table1_data_records.docx
(Times New Roman; title 10 pt bold; table text 9 pt; header bold; three-line table: thick top rule, thin rule under header, thin bottom rule;
footnote 8 pt). All values are read from source files; nothing is typed by hand.
Output: the tables directory set in config/config.ini (existing files are never overwritten)."""
import os, csv, json, hashlib, math
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = cfg.root + os.sep
RES = cfg.results_dir + os.sep
FIG = cfg.figdata_dir + os.sep
EV = cfg.qc_dir + os.sep
UP = cfg.data_root + os.sep
OUT = cfg.tables_dir + os.sep
os.makedirs(OUT, exist_ok=True)
S = cfg.samples
GSM = cfg.gsm
REPORT = []


def border(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side, sz in (("top", top), ("bottom", bottom)):
        if sz:
            e = OxmlElement("w:" + side)
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
            b.append(e)
    tcPr.append(b)


def run(par, text, size, bold=False, italic=False):
    """Text segments written as '^{...}' are set as Word superscript."""
    import re as _re
    parts = _re.split(r"(\^\{[^}]*\})", text)
    r = None
    for part in parts:
        if not part:
            continue
        sup = part.startswith("^{")
        r = par.add_run(part[2:-1] if sup else part)
        r.font.name = "Times New Roman"; r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
        r.font.size = Pt(size); r.bold = bold; r.italic = italic
        if sup:
            r.font.superscript = True
    return r


def build(fname, title, header, rows, footnotes, align_right_from=None, col_widths=None):
    path = OUT + fname
    if os.path.exists(path):
        print("[skip] %s exists — not overwritten" % fname)
        return path
    d = docx.Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = 1143000
    p = d.add_paragraph(); run(p, title, 10, bold=True)
    t = d.add_table(rows=len(rows) + 1, cols=len(header))
    t.alignment = 1
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""
        run(c.paragraphs[0], h, 9, bold=True)
        border(c, top=18, bottom=8)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]; c.text = ""
            run(c.paragraphs[0], str(v), 9)
            if align_right_from is not None and j >= align_right_from:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if i == len(rows):
                border(c, bottom=8)
    if col_widths:
        for j, w in enumerate(col_widths):
            for r_ in t.rows:
                r_.cells[j].width = w
    for fn in footnotes:
        q = d.add_paragraph(); run(q, fn, 8)
    d.save(path)
    print("[ok] %s (%d rows)" % (fname, len(rows)))
    return path


# ---------------- Table 1: data records with size and MD5 ----------------
md5 = {}
for f in ("raw_md5.txt", "processed_md5.txt"):
    for ln in open(UP + f, encoding="utf-8"):
        if ln.strip():
            h, pth = ln.split(None, 1); md5[pth.strip()] = h


def file_row(local, geo_name, desc, sample="", acc="GSE336465"):
    full = UP + local
    b = open(full, "rb").read()
    h = hashlib.md5(b).hexdigest()
    if h != md5[local]:
        raise SystemExit("STOP: MD5 mismatch for " + local)
    return [sample, acc, geo_name, desc, "{:,}".format(len(b)), h]


rows1 = []
for s in S:
    cond = "Mock infection" if s.startswith("MOCK") else "PR8 infection"
    rows1.append(file_row("raw/%s.txt" % s, "%s_%s.txt.gz" % (GSM[s], s), "Raw Feature Extraction file (%s, replicate %s)" % (cond, s[-1]), s, GSM[s]))
rows1.append(file_row("raw/Matrix.xlsx", "GSE336465_Matrix.xlsx", "Normalised matrix, 40,955 probes × 12 channels", "All"))
proc = [("mRNA_Quantity(flag_pass).xlsx", "m6A quantity and %modified, mRNA"), ("lncRNA_Quantity(flag_pass).xlsx", "m6A quantity and %modified, lncRNA"),
        ("others_Quantity(flag_pass).xlsx", "m6A quantity and %modified, other ncRNA"), ("mRNA_ExpressionLevel(flag_pass).xlsx", "Expression level, mRNA"),
        ("mRNA_Quantity(diffmethylated).xlsx", "Provider differential m6A table, mRNA"), ("lncRNA_Quantity(diffmethylated).xlsx", "Provider differential m6A table, lncRNA"),
        ("others_Quantity(diffmethylated).xlsx", "Provider differential m6A table, other ncRNA"), ("mRNA_ExpressionLevel(diffexpressed).xlsx", "Provider differential expression table, mRNA")]
for local, desc in proc:
    geo = "GSE336465_" + local.replace("(", "_").replace(")", "_")
    rows1.append(file_row("processed/" + local, geo, desc, "All"))
build("Table1_data_records.docx", "Table 1. Data records deposited in GEO (GSE336465).",
      ["Sample", "Accession", "File in GEO", "Content", "Size (bytes)", "MD5"], rows1,
      ["Sizes and MD5 checksums are those of the files as uploaded (the six Feature Extraction files are stored by GEO in gzip-compressed form). "
       "A README describing the files is also deposited with the Series. Provider tables use the MOCK-versus-PR8 convention (fold change = mock/PR8; see Usage Notes)."],
      align_right_from=4)
REPORT.append(("Table 1", len(rows1)))

# ---------------- Table 2: per-sample QC ----------------
t2 = list(csv.DictReader(open(RES + "table2_per_sample_qc.csv", encoding="utf-8-sig")))
rows2 = [[r["Sample"], r["Total RNA (µg)"], r["A260/A280"], r["A260/A230"], r["Positive spike-in in IP (% of input)"],
          r["Negative spike-in in IP (% of input)"], r["Sup (Cy3) cRNA yield (µg)"], r["Sup specific activity (pmol Cy3 per µg)"],
          r["IP (Cy5) cRNA yield (µg)"], r["IP specific activity (pmol Cy5 per µg)"]] for r in t2]
build("Table2_sample_quality_control.docx", "Table 2. RNA, immunoprecipitation and labelling quality control for each sample.",
      ["Sample", "Total RNA (µg)", "A260/A280", "A260/A230", "Positive spike-in in IP (% of input)", "Negative spike-in in IP (% of input)",
       "Sup cRNA (µg)", "Sup specific activity (pmol per µg)", "IP cRNA (µg)", "IP specific activity (pmol per µg)"], rows2,
      ["Total RNA: NanoDrop ND-1000 (15 µL per sample). Spike-in recovery: qPCR of the m6A-positive (CTA850) and m6A-negative (CTA650) spike-in RNAs, "
       "2^{−ΔCt}/(1 + 2^{−ΔCt}) × 100 with ΔCt = Ct(IP) − Ct(Sup). Labelling used 1 µg RNA per reaction; Sup was labelled with Cy3 and IP with Cy5; "
       "provider thresholds for two-colour hybridisation: yield ≥ 825 ng and specific activity ≥ 8.0 pmol per µg. Source: provider quality-control reports."],
      align_right_from=1)
REPORT.append(("Table 2", len(rows2)))

# ---------------- Table 3: array layout and calibration ----------------
t3 = list(csv.DictReader(open(RES + "table3_array_calibration.csv", encoding="utf-8-sig")))
img = {}
for r in t3:
    f = r["Scan image"]
    img.setdefault(f, len(img) + 1)
rows3 = [[r["Sample"], r["Array position"], "Image %d" % img[r["Scan image"]], r["IP constant (log2)"].replace("-", "−"), r["Sup constant (log2)"].replace("-", "−"),
          r["IP − Sup (log2)"].replace("-", "−"), r["Median log2(IP/Sup), deposited"].replace("-", "−"), r["Median %modified, deposited"]] for r in t3]
legend_img = "; ".join("Image %d = %s" % (v, k) for k, v in img.items())
build("Table3_array_layout_calibration.docx", "Table 3. Array layout and calibration constants.",
      ["Sample", "Array position", "Scan image", "IP constant (log2)", "Sup constant (log2)", "IP − Sup (log2)", "Median log2(IP/Sup)", "Median %modified"], rows3,
      ["Array position = slide_array as recorded by Feature Extraction (two 8×60K slides). Scan image files: %s. "
       "Constants are the per-channel offsets added by the provider's calibration (normalised log2 signal − log2 Feature Extraction signal; identical for all 40,955 probes of a channel). "
       "Medians are over all 40,955 probes of the deposited matrix." % legend_img],
      align_right_from=3)
REPORT.append(("Table 3", len(rows3)))

# ---------------- Table 4: interferon response ----------------
t4 = list(csv.DictReader(open(RES + "isg_transcript_table.csv", encoding="utf-8")))
def fdr(x):
    x = float(x); e = int(math.floor(math.log10(x))); m = x / 10 ** e
    if round(m, 1) >= 10:
        m, e = m / 10, e + 1
    return "%.1f × 10^{%s}" % (m, ("%d" % e).replace("-", "−"))
rows4 = []
for r in t4:
    rows4.append([r["gene"], r["transcript_id"], "%.2f" % float(r["sup_log2fc_pr8_vs_mock"]), fdr(r["sup_fdr"]),
                  "%.2f" % float(r["ip_log2fc_pr8_vs_mock"]), fdr(r["ip_fdr"])])
build("Table4_interferon_response.docx", "Table 4. Interferon and interferon-stimulated transcripts in PR8-infected and mock-infected cells.",
      ["Gene", "Transcript", "Expression log2 FC", "Expression FDR", "m6A quantity log2 FC", "m6A quantity FDR"], rows4,
      ["log2 FC = log2 fold change, PR8/mock; FDR = Benjamini–Hochberg-adjusted P value from limma on the deposited normalised matrix (three replicates per condition). "
       "Expression = Sup (Cy3) channel; m6A quantity = IP (Cy5) channel."],
      align_right_from=2)
REPORT.append(("Table 4", len(rows4)))
print("\n", REPORT)
