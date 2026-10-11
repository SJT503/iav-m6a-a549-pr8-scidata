# ---------------------------------------------------------------------------
# Paths and settings come from config/config.ini; nothing is machine-specific.
# ---------------------------------------------------------------------------
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _config import cfg

# -*- coding: utf-8 -*-
"""Build (1) Supplementary Table 1 (xlsx): probes whose %modified change is robust to normalisation (n = 210);
   (2) data for main-text Table 2 (per-sample RNA, spike-in and labelling QC) and Table 3 (array position, scan image,
   calibration constants), written as CSV for the manuscript build.
Sources: results/pctmod_robust_set_detected.csv (script 92), cache/matrix.csv, cache/probe_map_*.csv,
         provider_qc_evidence (NanoDrop, MeRIP qPCR), Agilent Labeling Report.pdf, FE file headers, results/norm_offsets_table.csv."""
import csv, os, math, statistics
import fitz, openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

RES = cfg.results_dir
CACHE = cfg.interim_dir
EV = cfg.qc_dir
RAW = cfg.raw_dir
OUT = cfg.tables_dir
os.makedirs(OUT, exist_ok=True)
S = cfg.samples
GSM = dict(zip(S, ["GSM9835854", "GSM9835855", "GSM9835856", "GSM9835857", "GSM9835858", "GSM9835859"]))

# ---------- (1) Supplementary Table 1 ----------
tid = {}
for k in ("mrna", "lncrna", "others"):
    for r in csv.DictReader(open(os.path.join(CACHE, "probe_map_%s.csv" % k), encoding="utf-8")):
        tid[r["probeID"]] = r["Transcript_ID"]
frac = {}
with open(os.path.join(CACHE, "matrix.csv"), encoding="utf-8-sig", newline="") as f:
    rd = csv.reader(f); h = next(rd)
    ci = [h.index("%s(IP,Normalized)" % s) for s in S]; cs = [h.index("%s(Sup,Normalized)" % s) for s in S]
    for r in rd:
        frac[r[0]] = [2 ** float(r[i]) / (2 ** float(r[i]) + 2 ** float(r[j])) for i, j in zip(ci, cs)]
rob = list(csv.DictReader(open(os.path.join(RES, "pctmod_robust_set_detected.csv"), encoding="utf-8")))
assert len(rob) == 210, len(rob)
rows, maxdev = [], 0.0
for r in rob:
    p = r["probeID"]; fr = frac[p]
    mm, mp = 100 * statistics.mean(fr[:3]), 100 * statistics.mean(fr[3:])
    d_si = 100 * float(r["dfrac_SI"])
    maxdev = max(maxdev, abs((mp - mm) - d_si))
    rows.append([p, tid[p], r["GeneSymbol"] or "", r["rna_class"].replace("other_ncRNA", "other ncRNA"),
                 round(mm, 2), round(mp, 2), round(d_si, 2), float(r["FDR_SI"]),
                 round(100 * float(r["dfrac_QNj"]), 2), float(r["FDR_QNj"]),
                 "higher in PR8" if r["direction"] == "higher_PR8" else "higher in mock"])
assert maxdev < 1e-6, maxdev
rows.sort(key=lambda x: (x[10] != "higher in PR8", -abs(x[6])))
n_up = sum(x[10] == "higher in PR8" for x in rows); n_dn = len(rows) - n_up

xl = os.path.join(OUT, "Supplementary_Table_1.xlsx")
if os.path.exists(xl):
    print("[skip] exists:", xl)
else:
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Supplementary Table 1"
    hdr = ["Probe ID", "Transcript ID", "Gene symbol", "RNA class",
           "Mean %modified, mock (deposited)", "Mean %modified, PR8 (deposited)",
           "Δ%modified, PR8 − mock (percentage points; deposited)", "FDR (deposited)",
           "Δ%modified, PR8 − mock (percentage points; quantile-normalised)", "FDR (quantile-normalised)", "Direction"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(name="Times New Roman", bold=True); c.alignment = Alignment(wrap_text=True, vertical="top")
    for x in rows:
        ws.append(x)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Times New Roman")
        row[7].number_format = "0.00E+00"; row[9].number_format = "0.00E+00"
        for k in (4, 5, 6, 8):
            row[k].number_format = "0.00"
    for k, w in enumerate([24, 18, 14, 12, 14, 14, 18, 12, 18, 12, 15], 1):
        ws.column_dimensions[get_column_letter(k)].width = w
    ws.freeze_panes = "A2"
    lg = wb.create_sheet("Legend")
    legend = [
        "Supplementary Table 1. Probes whose m6A methylation fraction (%modified) differs between PR8-infected and mock-infected A549 cells under both normalisations.",
        "%modified = 100 × IP/(IP + Sup), computed for each probe and array from the linear-scale (2^x) intensities of the deposited normalised matrix (Matrix.xlsx, GEO GSE336465).",
        "Testing was restricted to the 25,279 probes whose supernatant (Cy3) signal was well above background (Agilent Feature Extraction gIsWellAboveBG = 1) on all six arrays, because %modified is undefined when its denominator is at background.",
        "Differences were tested with limma (lmFit and eBayes; two-group design with mock as reference) on logit-transformed fractions; P values were adjusted by the Benjamini–Hochberg method within the 25,279 tested probes.",
        "Deposited: provider's calibration-spike-in normalisation as deposited in GEO. Quantile-normalised: the twelve channels (IP and Sup of the six arrays) were jointly quantile-normalised (limma normalizeBetweenArrays) before %modified was computed.",
        "A probe is listed if FDR < 0.05 and |Δ%modified| ≥ 5 percentage points in both analyses with the same sign (n = {}; {} higher in PR8, {} higher in mock). Each probe corresponds to one transcript; a gene can be represented by more than one probe.".format(len(rows), n_up, n_dn),
        "Rows are ordered by direction and then by the absolute deposited Δ%modified.",
    ]
    for t in legend:
        lg.append([t])
    lg.column_dimensions["A"].width = 160
    for row in lg.iter_rows():
        row[0].font = Font(name="Times New Roman"); row[0].alignment = Alignment(wrap_text=True, vertical="top")
    lg["A1"].font = Font(name="Times New Roman", bold=True)
    wb.save(xl)
    print("wrote", xl, "| rows", len(rows), "| higher PR8", n_up, "| higher mock", n_dn, "| max |Δ check| =", maxdev)

# ---------- (2) Table 2: per-sample QC ----------
nd = {r["sample"]: r for r in csv.DictReader(open(os.path.join(EV, "RNA_QC_nanoDrop_quantification.csv"), encoding="utf-8"))}
sp = {}
for r in csv.DictReader(open(os.path.join(EV, "MeRIP_QC_spikein_pct_input.csv"), encoding="utf-8")):
    sp.setdefault(r["sample"], {})[r["control_type"]] = r
txt = "".join(pg.get_text() for pg in fitz.open(os.path.join(EV, "Agilent Labeling Report.pdf")))
lab = {}
toks = [t.strip() for t in txt.split("\n") if t.strip()]
for i, t in enumerate(toks):
    if t.endswith(" sup") or t.endswith(" ip"):
        nm = t.replace("Mock-", "MOCK_").replace("PR8-", "PR8_").split()[0]
        frac_kind = "Sup" if t.endswith("sup") else "IP"
        dye, pmol, conc, sa, vol, tot = toks[i + 1:i + 7]
        lab[(nm, frac_kind)] = dict(dye=dye, sa=float(sa), total=float(tot))
assert len(lab) == 12, len(lab)
t2 = os.path.join(RES, "table2_per_sample_qc.csv")
with open(t2, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["Sample", "GEO accession", "Total RNA (µg)", "A260/A280", "A260/A230",
                "Positive spike-in in IP (% of input)", "Negative spike-in in IP (% of input)",
                "Sup (Cy3) cRNA yield (µg)", "Sup specific activity (pmol Cy3 per µg)",
                "IP (Cy5) cRNA yield (µg)", "IP specific activity (pmol Cy5 per µg)"])
    for s in S:
        w.writerow([s, GSM[s], "%.1f" % (float(nd[s]["total_rna_ng"]) / 1000), nd[s]["od260_280"], nd[s]["od260_230"],
                    "%.1f" % float(sp[s]["positive"]["pct_input"]), "%.2f" % float(sp[s]["negative"]["pct_input"]),
                    "%.2f" % lab[(s, "Sup")]["total"], "%.2f" % lab[(s, "Sup")]["sa"],
                    "%.2f" % lab[(s, "IP")]["total"], "%.2f" % lab[(s, "IP")]["sa"]])
print("wrote", t2)

# ---------- (3) Table 3: array position, scan image, calibration constants ----------
off = {r["sample"]: r for r in csv.DictReader(open(os.path.join(RES, "norm_offsets_table.csv"), encoding="utf-8"))}
const = {"MOCK_1": (-10.65155, -9.70681), "MOCK_2": (-10.17241, -9.50237), "MOCK_3": (-10.85019, -9.77775),
         "PR8_1": (-10.54372, -9.74515), "PR8_2": (-10.31455, -9.63554), "PR8_3": (-10.12483, -9.68526)}
t3 = os.path.join(RES, "table3_array_calibration.csv")
with open(t3, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["Sample", "Array position", "Scan image", "IP constant (log2)", "Sup constant (log2)", "IP − Sup (log2)",
                "Median log2(IP/Sup), deposited", "Median %modified, deposited"])
    for s in S:
        L = open(os.path.join(RAW, s + ".txt"), encoding="utf-8", errors="replace").read().split("\n")[:3]
        kv = dict(zip(L[1].split("\t"), L[2].split("\t")))
        img = os.path.basename(kv["FeatureExtractor_ScanFileName"].replace("\\", "/"))
        ip, su = const[s]
        assert abs((ip - su) - float(off[s]["calibration_factor_IP_minus_Sup"])) < 1e-3
        w.writerow([s, kv["FeatureExtractor_Barcode"], img, "%.3f" % ip, "%.3f" % su, "%.3f" % (ip - su),
                    "%.3f" % float(off[s]["deposited_median_log2ratio"]), "%.1f" % float(off[s]["deposited_median_pct_modified"])])
print("wrote", t3)
for p in (t2, t3):
    print("\n" + open(p, encoding="utf-8-sig").read())
